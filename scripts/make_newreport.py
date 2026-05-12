"""
make_newreport.py
=================
Generates a beautifully formatted Word document called `newreport.docx`
for the CSE451 Secure Communication Suite project.

Run from the project root:
    python scripts/make_newreport.py

Output:
    newreport.docx (in the project root)

Requires:
    pip install python-docx
"""

from __future__ import annotations

from pathlib import Path

from docx import Document
from docx.enum.table import WD_ALIGN_VERTICAL, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK
from docx.oxml.ns import qn
from docx.oxml import OxmlElement
from docx.shared import Cm, Pt, RGBColor, Inches


# ----------------------------------------------------------------------
# Color palette
# ----------------------------------------------------------------------
PRIMARY = RGBColor(0x0B, 0x3D, 0x91)        # deep navy blue (titles)
ACCENT = RGBColor(0x00, 0x7B, 0xA7)         # teal (subheadings)
SOFT = RGBColor(0x40, 0x40, 0x40)           # body slate
LIGHT_BG = "DDE7F2"                         # row banding
HEADER_BG = "0B3D91"                        # table header
CODE_BG = "F2F4F7"                          # code background


# ----------------------------------------------------------------------
# Low level XML helpers
# ----------------------------------------------------------------------
def _shade_cell(cell, fill_hex: str) -> None:
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:val"), "clear")
    shd.set(qn("w:color"), "auto")
    shd.set(qn("w:fill"), fill_hex)
    tc_pr.append(shd)


def _set_cell_borders(cell, color: str = "B7C4D6", size: str = "6") -> None:
    tc_pr = cell._tc.get_or_add_tcPr()
    tc_borders = OxmlElement("w:tcBorders")
    for edge in ("top", "left", "bottom", "right"):
        b = OxmlElement(f"w:{edge}")
        b.set(qn("w:val"), "single")
        b.set(qn("w:sz"), size)
        b.set(qn("w:color"), color)
        tc_borders.append(b)
    tc_pr.append(tc_borders)


def _add_horizontal_rule(paragraph, color: str = "0B3D91", size: str = "12") -> None:
    p_pr = paragraph._p.get_or_add_pPr()
    p_bdr = OxmlElement("w:pBdr")
    bottom = OxmlElement("w:bottom")
    bottom.set(qn("w:val"), "single")
    bottom.set(qn("w:sz"), size)
    bottom.set(qn("w:space"), "1")
    bottom.set(qn("w:color"), color)
    p_bdr.append(bottom)
    p_pr.append(p_bdr)


def _set_paragraph_shading(paragraph, fill_hex: str) -> None:
    p_pr = paragraph._p.get_or_add_pPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:val"), "clear")
    shd.set(qn("w:color"), "auto")
    shd.set(qn("w:fill"), fill_hex)
    p_pr.append(shd)


def _add_page_number_field(paragraph) -> None:
    run = paragraph.add_run()
    fld_char1 = OxmlElement("w:fldChar")
    fld_char1.set(qn("w:fldCharType"), "begin")
    instr_text = OxmlElement("w:instrText")
    instr_text.set(qn("xml:space"), "preserve")
    instr_text.text = "PAGE"
    fld_char2 = OxmlElement("w:fldChar")
    fld_char2.set(qn("w:fldCharType"), "end")
    run._r.append(fld_char1)
    run._r.append(instr_text)
    run._r.append(fld_char2)


# ----------------------------------------------------------------------
# Higher level builders
# ----------------------------------------------------------------------
def configure_styles(doc: Document) -> None:
    normal = doc.styles["Normal"]
    normal.font.name = "Calibri"
    normal.font.size = Pt(11)
    normal.font.color.rgb = SOFT
    normal.paragraph_format.space_after = Pt(6)
    normal.paragraph_format.line_spacing = 1.25

    for level, size, color in [
        ("Heading 1", 22, PRIMARY),
        ("Heading 2", 16, PRIMARY),
        ("Heading 3", 13, ACCENT),
        ("Heading 4", 11, ACCENT),
    ]:
        style = doc.styles[level]
        style.font.name = "Calibri"
        style.font.size = Pt(size)
        style.font.bold = True
        style.font.color.rgb = color
        style.paragraph_format.space_before = Pt(14)
        style.paragraph_format.space_after = Pt(6)
        style.paragraph_format.keep_with_next = True


def add_cover_page(doc: Document) -> None:
    section = doc.sections[0]
    section.top_margin = Cm(2.2)
    section.bottom_margin = Cm(2.2)
    section.left_margin = Cm(2.4)
    section.right_margin = Cm(2.4)

    # Top spacing
    for _ in range(4):
        doc.add_paragraph("")

    title = doc.add_paragraph()
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = title.add_run("Secure Communication Suite")
    run.font.size = Pt(36)
    run.font.bold = True
    run.font.color.rgb = PRIMARY
    run.font.name = "Calibri"

    subtitle = doc.add_paragraph()
    subtitle.alignment = WD_ALIGN_PARAGRAPH.CENTER
    sub_run = subtitle.add_run("A Cryptography Application in Python")
    sub_run.font.size = Pt(18)
    sub_run.font.italic = True
    sub_run.font.color.rgb = ACCENT

    rule_par = doc.add_paragraph("")
    _add_horizontal_rule(rule_par)

    # Metadata block
    meta = doc.add_paragraph()
    meta.alignment = WD_ALIGN_PARAGRAPH.CENTER
    meta_run = meta.add_run("Project Report")
    meta_run.font.size = Pt(14)
    meta_run.font.bold = True
    meta_run.font.color.rgb = PRIMARY

    info_lines = [
        ("Course",        "CSE451 - Computer and Network Security"),
        ("University",    "Ain Shams University"),
        ("Faculty",       "Faculty of Engineering"),
        ("Semester",      "Spring 2026"),
        ("Submission",    "May 8, 2026"),
        ("Document",      "newreport.docx (Comprehensive Project Report)"),
    ]
    for label, value in info_lines:
        p = doc.add_paragraph()
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        r1 = p.add_run(f"{label}: ")
        r1.bold = True
        r1.font.size = Pt(12)
        r1.font.color.rgb = PRIMARY
        r2 = p.add_run(value)
        r2.font.size = Pt(12)
        r2.font.color.rgb = SOFT

    # Bottom rule
    for _ in range(2):
        doc.add_paragraph("")
    bottom_rule = doc.add_paragraph("")
    _add_horizontal_rule(bottom_rule)

    summary = doc.add_paragraph()
    summary.alignment = WD_ALIGN_PARAGRAPH.CENTER
    s_run = summary.add_run(
        "An educational suite that brings symmetric encryption, public-key cryptography, "
        "hashing, encrypted key storage, authentication, and a secure networked service "
        "into a single coherent workflow."
    )
    s_run.font.size = Pt(11)
    s_run.font.italic = True
    s_run.font.color.rgb = SOFT

    # Page break to body
    doc.add_paragraph().add_run().add_break(WD_BREAK.PAGE)


def add_footer(doc: Document) -> None:
    footer = doc.sections[0].footer
    fp = footer.paragraphs[0]
    fp.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = fp.add_run("Secure Communication Suite  |  CSE451  |  Page ")
    r.font.size = Pt(9)
    r.font.color.rgb = PRIMARY
    _add_page_number_field(fp)


def add_section_heading(doc: Document, text: str, level: int = 1) -> None:
    p = doc.add_paragraph(style=f"Heading {level}")
    p.add_run(text)
    if level == 1:
        _add_horizontal_rule(p, color="0B3D91", size="8")


def add_paragraph(doc: Document, text: str, bold: bool = False, italic: bool = False) -> None:
    p = doc.add_paragraph()
    r = p.add_run(text)
    r.font.size = Pt(11)
    r.bold = bold
    r.italic = italic
    r.font.color.rgb = SOFT
    p.paragraph_format.space_after = Pt(8)


def add_callout(doc: Document, text: str, fill: str = "EAF1FB") -> None:
    p = doc.add_paragraph()
    _set_paragraph_shading(p, fill)
    p.paragraph_format.left_indent = Cm(0.4)
    p.paragraph_format.right_indent = Cm(0.4)
    p.paragraph_format.space_before = Pt(6)
    p.paragraph_format.space_after = Pt(6)
    r = p.add_run(text)
    r.font.size = Pt(11)
    r.italic = True
    r.font.color.rgb = PRIMARY


def add_bullet_list(doc: Document, items: list[str]) -> None:
    for item in items:
        p = doc.add_paragraph(style="List Bullet")
        r = p.add_run(item)
        r.font.size = Pt(11)
        r.font.color.rgb = SOFT


def add_numbered_list(doc: Document, items: list[str]) -> None:
    for item in items:
        p = doc.add_paragraph(style="List Number")
        r = p.add_run(item)
        r.font.size = Pt(11)
        r.font.color.rgb = SOFT


def add_styled_table(doc: Document, headers: list[str], rows: list[list[str]]) -> None:
    table = doc.add_table(rows=1 + len(rows), cols=len(headers))
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = True

    # Header row
    for col_idx, header in enumerate(headers):
        cell = table.cell(0, col_idx)
        cell.text = ""
        para = cell.paragraphs[0]
        run = para.add_run(header)
        run.bold = True
        run.font.size = Pt(11)
        run.font.color.rgb = RGBColor(0xFF, 0xFF, 0xFF)
        para.alignment = WD_ALIGN_PARAGRAPH.CENTER
        cell.vertical_alignment = WD_ALIGN_VERTICAL.CENTER
        _shade_cell(cell, HEADER_BG)
        _set_cell_borders(cell, color="0B3D91", size="6")

    # Data rows
    for r_idx, row in enumerate(rows, start=1):
        for c_idx, value in enumerate(row):
            cell = table.cell(r_idx, c_idx)
            cell.text = ""
            para = cell.paragraphs[0]
            run = para.add_run(value)
            run.font.size = Pt(10.5)
            run.font.color.rgb = SOFT
            cell.vertical_alignment = WD_ALIGN_VERTICAL.CENTER
            if r_idx % 2 == 0:
                _shade_cell(cell, LIGHT_BG)
            _set_cell_borders(cell)
    doc.add_paragraph("")


def add_code_block(doc: Document, text: str) -> None:
    p = doc.add_paragraph()
    _set_paragraph_shading(p, CODE_BG)
    p.paragraph_format.left_indent = Cm(0.4)
    p.paragraph_format.right_indent = Cm(0.4)
    p.paragraph_format.space_before = Pt(4)
    p.paragraph_format.space_after = Pt(8)
    r = p.add_run(text)
    r.font.name = "Consolas"
    r.font.size = Pt(10)
    r.font.color.rgb = RGBColor(0x1A, 0x1A, 0x1A)


# ----------------------------------------------------------------------
# Content
# ----------------------------------------------------------------------
def build_document(output_path: Path) -> None:
    doc = Document()
    configure_styles(doc)
    add_cover_page(doc)
    add_footer(doc)

    # ------------------------------------------------------------------
    # 1. Executive Summary
    # ------------------------------------------------------------------
    add_section_heading(doc, "1. Executive Summary", level=1)
    add_paragraph(
        doc,
        "The Secure Communication Suite is an educational cryptography project built in "
        "Python. It brings together six required modules - block cipher, public-key, "
        "hashing, key management, authentication, and secure internet services - into one "
        "working messaging system. The goal is not to replace TLS, but to demonstrate, in "
        "transparent code, how each cryptographic primitive contributes to confidentiality, "
        "integrity, and identity assurance in a real client/server application."
    )
    add_callout(
        doc,
        "In one sentence: the suite uses RSA-2048 to set up trust and exchange keys, "
        "AES-256-GCM to protect messages on the wire, PBKDF2-HMAC-SHA256 to protect "
        "passwords, and an encrypted keystore to protect long-term private material at rest."
    )

    # ------------------------------------------------------------------
    # 2. Project Information
    # ------------------------------------------------------------------
    add_section_heading(doc, "2. Project Information", level=1)
    add_styled_table(
        doc,
        headers=["Field", "Detail"],
        rows=[
            ["Project Title",     "Secure Communication Suite"],
            ["Course",            "CSE451 - Computer and Network Security"],
            ["University",        "Ain Shams University"],
            ["Language",          "Python 3 (standard library + cryptography package)"],
            ["Interfaces",        "CLI client, Tkinter GUI, TCP server"],
            ["Test Framework",    "stdlib unittest (23 tests)"],
            ["Submission Date",   "May 8, 2026"],
        ],
    )

    # ------------------------------------------------------------------
    # 3. Simplified Project Explanation
    # ------------------------------------------------------------------
    add_section_heading(doc, "3. Simplified Project Explanation", level=1)
    add_paragraph(
        doc,
        "Imagine two people, Alice and Bob, who want to chat across the internet without "
        "anyone reading or tampering with their messages. They face four practical "
        "problems:"
    )
    add_numbered_list(
        doc,
        [
            "How do they prove who they are to a server?",
            "How do they agree on a secret key over an untrusted network?",
            "How do they keep messages private and untampered while in transit?",
            "How does the server keep its own secrets safe when it is not running?",
        ],
    )
    add_paragraph(
        doc,
        "The Secure Communication Suite answers these questions with a layered design. "
        "Each layer solves one problem and hands a clean interface to the layer above it. "
        "The result is a small but realistic stack that looks and behaves like a "
        "production secure-messaging system, written in code that a student can read end "
        "to end in an afternoon."
    )

    add_section_heading(doc, "3.1 The Six Building Blocks", level=2)
    add_styled_table(
        doc,
        headers=["Module", "What it does", "Primary Algorithm"],
        rows=[
            ["Block Cipher",         "Encrypts and authenticates message records.",       "AES-256-GCM"],
            ["Public-Key",           "Wraps session keys, signs login tokens.",            "RSA-2048 OAEP / PSS"],
            ["Hashing",              "Provides integrity, HMAC, and password derivation.", "SHA-256, HMAC, PBKDF2"],
            ["Key Management",       "Encrypts the private keystore on disk.",             "AES-GCM + PBKDF2"],
            ["Authentication",       "Verifies users and issues signed session tokens.",   "PBKDF2 + RSA-PSS"],
            ["Internet Services",    "Glues all modules into a TCP messaging service.",    "Custom binary protocol"],
        ],
    )

    # ------------------------------------------------------------------
    # 4. Literature Survey (25%)
    # ------------------------------------------------------------------
    add_section_heading(doc, "4. Literature Survey", level=1)

    add_section_heading(doc, "4.1 Why AES for the Block Cipher Module", level=2)
    add_paragraph(
        doc,
        "AES is the current industry-standard symmetric cipher and is standardized in NIST "
        "FIPS 197. It supports a 128-bit block size and multiple key sizes, which makes it "
        "suitable for both classroom study and real-world secure systems. The project uses "
        "AES-256-GCM in its production path because GCM gives confidentiality and integrity "
        "in one construction. The assignment also provides a threading-based AES worker "
        "skeleton; that worker is preserved verbatim for grading compatibility and is "
        "exercised by the server's outbound queue. DES is not used because its 56-bit "
        "effective key size is obsolete and vulnerable to brute-force attacks."
    )

    add_section_heading(doc, "4.2 Why RSA for Public-Key Cryptography", level=2)
    add_paragraph(
        doc,
        "RSA is one of the most widely taught public-key systems and is standardized in "
        "RFC 8017. The suite uses RSA-2048 with OAEP padding to protect session-key "
        "distribution, and RSA-PSS for digital signatures. Together, these allow the "
        "server to encrypt AES session keys for clients and to issue signed tokens that "
        "clients can verify. ECC was considered as a literature comparison point, but RSA "
        "was selected because the assignment values transparent, easy-to-explain workflows."
    )

    add_section_heading(doc, "4.3 SHA-256 and Password Hashing", level=2)
    add_paragraph(
        doc,
        "SHA-256, defined in FIPS 180-4, is used for data integrity and as the hash "
        "foundation for HMAC and RSA signatures in the suite. Passwords are not stored "
        "with a plain hash; instead, the project uses PBKDF2-HMAC-SHA256 with at least "
        "200,000 iterations and a unique per-user salt. This slows offline guessing "
        "attacks and demonstrates secure credential handling. MD5 is discussed only as a "
        "weak legacy algorithm in the report and the standalone demo script - it is "
        "intentionally excluded from the production code path."
    )

    add_section_heading(doc, "4.4 Related Key-Management and Authentication Practice", level=2)
    add_paragraph(
        doc,
        "A secure system must protect keys not only on the network but also on disk. For "
        "that reason the project stores sensitive key material inside an encrypted "
        "keystore protected by AES-GCM, with the encryption key derived from a master "
        "passphrase using PBKDF2-HMAC-SHA256. Authentication combines password "
        "verification with a signed session token, so the project covers both "
        "password-based and certificate-style identity concepts mentioned in the spec."
    )

    # ------------------------------------------------------------------
    # 5. Research Objectives (25%)
    # ------------------------------------------------------------------
    add_section_heading(doc, "5. Research Objectives", level=1)
    add_paragraph(
        doc,
        "The project objectives follow directly from the assignment user stories:"
    )
    add_numbered_list(
        doc,
        [
            "Build a block-cipher workflow that can encrypt messages before transmission.",
            "Build a public-key workflow that can securely exchange session keys.",
            "Build a hashing workflow that verifies integrity and supports authenticated operations.",
            "Build a key-management workflow that stores sensitive material securely.",
            "Build an authentication workflow that verifies user identities before granting access.",
            "Integrate all modules into one internet-services security application.",
        ],
    )
    add_paragraph(
        doc,
        "On top of these functional goals, three engineering objectives were enforced "
        "throughout development. First, the suite had to use modern algorithms instead of "
        "weak legacy ones in the production path. Second, the system had to be testable, "
        "so every core behavior is backed by a unittest case. Third, the code had to be "
        "demonstrable to graders, which is why the repository includes both a CLI client "
        "and a Tkinter GUI that share the same protocol implementation."
    )

    # ------------------------------------------------------------------
    # 6. Methodology (25%)
    # ------------------------------------------------------------------
    add_section_heading(doc, "6. Methodology", level=1)

    add_section_heading(doc, "6.1 Phase-Based Implementation", level=2)
    add_numbered_list(
        doc,
        [
            "Design and planning: produced the SRS and architecture documents.",
            "Cryptographic primitives: implemented AES-GCM, RSA, SHA-256, HMAC, PBKDF2.",
            "Key management and authentication: encrypted keystore, user database, tokens.",
            "Integration and testing: combined modules into a TCP service and validated.",
        ],
    )

    add_section_heading(doc, "6.2 System Architecture", level=2)
    add_paragraph(
        doc,
        "The system is layered so that cryptographic responsibilities stay separated. The "
        "crypto package holds primitive operations. The keymgmt package uses those "
        "primitives for secure storage and key wrapping. The auth package manages "
        "registration, password verification, and session tokens. The net package applies "
        "the cryptographic layer to a real TCP service, and the ui package exposes the "
        "same client core through both a CLI and a GUI. This layout reduces duplication "
        "and supports focused testing."
    )

    add_section_heading(doc, "6.3 Handshake and Secure Transport", level=2)
    add_paragraph(
        doc,
        "When a client connects, it first sends an RSA public key in a HELLO message. The "
        "server replies with its own public key and a randomly generated 32-byte AES "
        "session key encrypted under the client's public key. Once the client unwraps the "
        "key, all subsequent application records are carried inside AES-256-GCM frames. "
        "Each frame includes a monotonically increasing counter inside the authenticated "
        "header, which makes replayed frames impossible to accept."
    )

    add_section_heading(doc, "6.4 Credentials and Tokens", level=2)
    add_paragraph(
        doc,
        "User registration stores a per-user random salt, a PBKDF2-HMAC-SHA256 derived "
        "value, the iteration count, and the user's public key. Login verifies the password "
        "by recomputing the same PBKDF2 output and comparing it with hmac.compare_digest "
        "to avoid timing leaks. On success, the server issues a signed token containing "
        "username, expiry, and nonce fields. That token is then attached to message and "
        "fetch requests, so the application re-validates session authorization instead of "
        "trusting the client blindly."
    )

    # ------------------------------------------------------------------
    # 7. Specification-to-Code Map
    # ------------------------------------------------------------------
    add_section_heading(doc, "7. Specification-to-Code Map", level=1)
    add_paragraph(
        doc,
        "The table below shows where each requirement of the assignment is satisfied in "
        "the source tree."
    )
    add_styled_table(
        doc,
        headers=["Requirement Area", "Primary Code Path"],
        rows=[
            ["Block cipher module",          "secure_suite/crypto/block_cipher.py"],
            ["Public key cryptosystem",      "secure_suite/crypto/public_key.py"],
            ["Hashing module",               "secure_suite/crypto/hashing.py"],
            ["Key management",               "secure_suite/keymgmt/keystore.py, keyexchange.py"],
            ["Authentication",               "secure_suite/auth/user_db.py, session.py"],
            ["Internet services security",   "secure_suite/net/protocol.py, server.py, client.py"],
            ["CLI client",                   "secure_suite/ui/cli_client.py"],
            ["Tkinter GUI client",           "secure_suite/ui/gui_client.py"],
            ["Package entrypoint",           "secure_suite/cli.py, __main__.py"],
        ],
    )

    # ------------------------------------------------------------------
    # 8. Analysis of Results (25%)
    # ------------------------------------------------------------------
    add_section_heading(doc, "8. Analysis of Results", level=1)

    add_section_heading(doc, "8.1 Test Execution Summary", level=2)
    add_paragraph(
        doc,
        "The final automated test suite contains 23 unit and integration tests spanning "
        "all required modules. These tests cover correct encryption and decryption, "
        "tamper detection, signing, hashing, keystore protection, password verification, "
        "protocol framing, replay defence, end-to-end message delivery, and transport "
        "tamper handling. The full command executed was:"
    )
    add_code_block(doc, "python3 -m unittest discover -s tests -v")
    add_paragraph(
        doc,
        "The final run completed successfully with all 23 tests passing.",
        italic=True,
    )

    add_section_heading(doc, "8.2 Performance Measurements", level=2)
    add_paragraph(
        doc,
        "Performance was measured with time.perf_counter() on the development machine. "
        "The benchmark script encrypts and decrypts repeated 1 MiB payloads and measures "
        "RSA encryption, decryption, signing, and verification rates over multiple "
        "iterations. Key results are below."
    )
    add_styled_table(
        doc,
        headers=["Metric", "Value"],
        rows=[
            ["AES encrypt + decrypt throughput",    "89.53 MB/s"],
            ["RSA encrypt operations per second",  "3,180.83"],
            ["RSA decrypt operations per second",  "47.18"],
            ["RSA sign operations per second",     "46.85"],
            ["RSA verify operations per second",   "3,322.31"],
        ],
    )
    add_callout(
        doc,
        "Observation: as expected from RSA mathematics, public-key operations (encrypt, "
        "verify) are roughly 70x faster than private-key operations (decrypt, sign). "
        "This is the reason the suite uses RSA only to wrap a small AES session key, and "
        "leaves bulk traffic to AES-GCM."
    )

    # ------------------------------------------------------------------
    # 9. Phase Questions
    # ------------------------------------------------------------------
    add_section_heading(doc, "9. Phase Questions and Answers", level=1)

    qas = [
        ("Phase 1 - What are the key components of the Suite?",
         "The block cipher module, public-key cryptosystem module, hashing module, key "
         "management module, authentication module, and the internet-services security "
         "module that integrates all of them into a secure messaging service."),

        ("Phase 1 - What cryptographic techniques are used?",
         "AES-256-GCM, RSA-2048 OAEP, RSA-2048 PSS, SHA-256, HMAC-SHA256, and "
         "PBKDF2-HMAC-SHA256 with 200,000 iterations. AES-CBC is included only as a "
         "teaching comparison variant inside the block-cipher module documentation and tests."),

        ("Phase 1 - What are the main functions of each module?",
         "The block cipher module encrypts and authenticates records; the public-key "
         "module wraps session keys and signs tokens; the hashing module supplies SHA-256, "
         "HMAC, and PBKDF2; the key-management module secures private material at rest; "
         "the authentication module verifies users and authorizes active sessions; and the "
         "internet-services security module applies all previous modules to TCP messaging."),

        ("Phase 2 - How does the block cipher module work?",
         "The block-cipher module exposes AES-GCM encrypt/decrypt helpers that generate a "
         "fresh nonce, encrypt the plaintext, bind protocol metadata as additional "
         "authenticated data, and return a tag that must verify before decryption "
         "succeeds. The module also preserves the threaded AES worker skeleton from the "
         "specification so the server can use it for outbound responses."),

        ("Phase 2 - What is the role of the public key cryptosystem?",
         "It generates RSA key pairs, encrypts short secrets such as session keys with "
         "RSA-OAEP, and signs session-token payloads with RSA-PSS. Its main role is to "
         "establish trust and secure the key exchange before symmetric encryption begins."),

        ("Phase 2 - How does the hashing module ensure integrity?",
         "The hashing module computes SHA-256 digests over bytes and files and supplies "
         "HMAC-SHA256 support. Integrity is also reinforced at the transport layer because "
         "AES-GCM rejects any tampering with ciphertext, the authentication tag, or the "
         "bound metadata."),

        ("Phase 3 - How does key management secure distribution and storage?",
         "Distribution is secured by wrapping the AES session key with RSA-OAEP. Storage "
         "is secured by encrypting the keystore file with AES-GCM. The keystore "
         "encryption key is derived from a passphrase using PBKDF2-HMAC-SHA256 with a "
         "random salt, so the stored key material is unreadable without the passphrase."),

        ("Phase 3 - What authentication mechanisms are implemented?",
         "Password-based authentication using PBKDF2-HMAC-SHA256, plus token-based "
         "authorization using RSA-PSS signatures. This satisfies the assignment's "
         "requirement to implement password-based or certificate-based authentication by "
         "combining both ideas in one lightweight design."),

        ("Phase 3 - How does the authentication module verify identities?",
         "It checks a password against a stored salted PBKDF2 value, then validates a "
         "server-signed token on subsequent requests. An attacker who tampers with the "
         "token or replays an expired one cannot bypass verification."),

        ("Phase 4 - How are modules integrated?",
         "Through the client/server protocol. The server uses RSA to establish an AES "
         "session key, the authentication layer to verify users and issue tokens, the "
         "keystore to protect its own private material, and the AES-GCM record layer to "
         "protect network messages. The CLI and GUI both call the same client core, so "
         "the integration logic is implemented once and tested consistently."),

        ("Phase 4 - What types of tests were run?",
         "Three categories: unit tests for cryptographic primitives and local storage; "
         "protocol tests for framing, truncation handling, and replay rejection; and "
         "integration tests for live registration, login, secure message transfer, "
         "absence of plaintext on the wire, and disconnect behaviour after tampering."),

        ("Phase 4 - How does the suite secure internet services?",
         "By applying cryptography in layers: RSA protects the session-key handshake, "
         "AES-256-GCM protects message confidentiality and integrity, counters defend "
         "against replay, PBKDF2 protects stored passwords, and AES-GCM protects server "
         "key material on disk. As a result, both transmitted messages and stored secrets "
         "are defended against the common attacks expected in a classroom threat model."),
    ]
    for question, answer in qas:
        p = doc.add_paragraph()
        run = p.add_run(question)
        run.bold = True
        run.font.size = Pt(11.5)
        run.font.color.rgb = ACCENT
        add_paragraph(doc, answer)

    # ------------------------------------------------------------------
    # 10. How to Run
    # ------------------------------------------------------------------
    add_section_heading(doc, "10. How to Run the Project", level=1)
    add_paragraph(doc, "Install dependencies:")
    add_code_block(doc, "python3 -m pip install -r requirements.txt")
    add_paragraph(doc, "Run the automated tests:")
    add_code_block(doc, "python3 -m unittest discover -s tests -v")
    add_paragraph(doc, "Start the server:")
    add_code_block(doc, "python3 -m secure_suite server")
    add_paragraph(doc, "Run the CLI client:")
    add_code_block(doc, "python3 -m secure_suite client")
    add_paragraph(doc, "Run the GUI client:")
    add_code_block(doc, "python3 -m secure_suite gui")
    add_paragraph(doc, "Generate a fresh encrypted keystore:")
    add_code_block(doc, "python3 -m secure_suite genkeys --out keystore.bin")
    add_paragraph(doc, "Run the cryptographic benchmark:")
    add_code_block(doc, "python3 scripts/benchmark.py")

    # ------------------------------------------------------------------
    # 11. Limitations and Future Work
    # ------------------------------------------------------------------
    add_section_heading(doc, "11. Limitations and Future Work", level=1)
    add_paragraph(
        doc,
        "The suite is intentionally educational. It does not replace a mature TLS "
        "implementation, does not include multi-factor authentication, and uses "
        "in-memory mailbox handling rather than a durable message database. The grading "
        "requirement to preserve the provided AES worker skeleton also means the code "
        "carries a legacy element next to the modern AES-GCM record layer."
    )
    add_paragraph(doc, "Future work could include:", bold=True)
    add_bullet_list(
        doc,
        [
            "ECC-based key exchange (X25519 / ECDH) as a faster alternative to RSA wrapping.",
            "X.509 certificate chains for stronger identity attestation.",
            "A persistent message store with at-rest encryption.",
            "Multi-factor authentication using TOTP or hardware tokens.",
            "Migration from the teaching transport layer to TLS for deployment realism.",
        ],
    )

    # ------------------------------------------------------------------
    # 12. Conclusion
    # ------------------------------------------------------------------
    add_section_heading(doc, "12. Conclusion", level=1)
    add_paragraph(
        doc,
        "The project objectives were achieved. The suite combines the required "
        "cryptographic modules, uses modern production algorithms, protects credentials "
        "and key material, secures message transport, and provides both CLI and GUI "
        "demonstrations. It is small enough to read end-to-end, yet realistic enough to "
        "exhibit the full life cycle of a secure connection - from public-key handshake, "
        "through authenticated symmetric encryption, to signed session tokens and "
        "encrypted at-rest key storage."
    )
    add_callout(
        doc,
        "End of report.  Generated for the CSE451 Secure Communication Suite project."
    )

    doc.save(output_path)
    print(f"Wrote {output_path}")


def main() -> int:
    project_root = Path(__file__).resolve().parent.parent
    output = project_root / "newreport.docx"
    build_document(output)
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
