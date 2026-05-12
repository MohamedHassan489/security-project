"""Render markdown project deliverables without pandoc."""

from __future__ import annotations

import argparse
from pathlib import Path

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import Inches, Pt
from fpdf import FPDF
from pptx import Presentation
from pptx.enum.text import PP_ALIGN
from pptx.util import Inches as PptInches
from pptx.util import Pt as PptPt


def normalize_pdf_text(text: str) -> str:
    replacements = {
        "•": "-",
        "–": "-",
        "—": "-",
        "’": "'",
        "“": '"',
        "”": '"',
        "≈": "~",
    }
    for source, target in replacements.items():
        text = text.replace(source, target)
    return text.encode("latin-1", "replace").decode("latin-1")


def pdf_multi_cell(pdf: FPDF, height: float, text: str) -> None:
    pdf.set_x(pdf.l_margin)
    pdf.multi_cell(0, height, text)


def parse_markdown_tables(lines: list[str], start_index: int) -> tuple[list[list[str]], int]:
    table_lines: list[str] = []
    index = start_index
    while index < len(lines) and lines[index].strip().startswith("|"):
        table_lines.append(lines[index].rstrip())
        index += 1
    rows = []
    for line in table_lines:
        cells = [cell.strip() for cell in line.strip().strip("|").split("|")]
        rows.append(cells)
    if len(rows) >= 2 and all(cell.replace("-", "").replace(":", "") == "" for cell in rows[1]):
        rows.pop(1)
    return rows, index


def is_numbered_line(stripped: str) -> bool:
    return len(stripped) > 3 and stripped[0].isdigit() and stripped[1:3] == ". "


def is_special_line(stripped: str) -> bool:
    return (
        not stripped
        or stripped.startswith("#")
        or stripped.startswith("- ")
        or is_numbered_line(stripped)
        or stripped.startswith("|")
        or stripped.startswith("```")
    )


def add_docx_table(document: Document, rows: list[list[str]]) -> None:
    if not rows:
        return
    table = document.add_table(rows=len(rows), cols=len(rows[0]))
    table.style = "Table Grid"
    for row_index, row in enumerate(rows):
        for col_index, value in enumerate(row):
            table.cell(row_index, col_index).text = value
    document.add_paragraph("")


def render_docx(markdown_path: Path, output_path: Path) -> None:
    lines = markdown_path.read_text(encoding="utf-8").splitlines()
    document = Document()
    normal_style = document.styles["Normal"]
    normal_style.font.name = "Times New Roman"
    normal_style.font.size = Pt(11)

    index = 0
    in_code_block = False
    while index < len(lines):
        line = lines[index].rstrip()
        stripped = line.strip()
        if stripped.startswith("```"):
            in_code_block = not in_code_block
            index += 1
            continue
        if in_code_block:
            paragraph = document.add_paragraph()
            paragraph.style = document.styles["No Spacing"]
            run = paragraph.add_run(line)
            run.font.name = "Courier New"
            run.font.size = Pt(9)
            index += 1
            continue
        if not stripped:
            document.add_paragraph("")
            index += 1
            continue
        if stripped.startswith("|"):
            rows, index = parse_markdown_tables(lines, index)
            add_docx_table(document, rows)
            continue
        if stripped.startswith("#"):
            level = len(stripped) - len(stripped.lstrip("#"))
            text = stripped[level:].strip()
            paragraph = document.add_paragraph(style=f"Heading {min(level, 4)}")
            paragraph.add_run(text)
            index += 1
            continue
        if stripped.startswith("- "):
            paragraph = document.add_paragraph(style="List Bullet")
            paragraph.add_run(stripped[2:].strip())
            index += 1
            continue
        if is_numbered_line(stripped):
            while index < len(lines) and is_numbered_line(lines[index].strip()):
                document.add_paragraph(lines[index].strip()[3:].strip(), style="List Number")
                index += 1
            continue
        paragraph_lines = [stripped]
        index += 1
        while index < len(lines):
            candidate = lines[index].strip()
            if is_special_line(candidate):
                break
            paragraph_lines.append(candidate)
            index += 1
        paragraph = document.add_paragraph(" ".join(paragraph_lines))
        if stripped.startswith("**") and stripped.endswith("**") and len(paragraph.runs) == 1:
            paragraph.runs[0].bold = True
        continue
    document.save(output_path)


def render_pdf(markdown_path: Path, output_path: Path) -> None:
    lines = markdown_path.read_text(encoding="utf-8").splitlines()
    pdf = FPDF()
    pdf.set_auto_page_break(auto=True, margin=15)
    pdf.add_page()
    pdf.set_font("Helvetica", size=11)
    index = 0
    in_code_block = False
    while index < len(lines):
        line = lines[index].rstrip()
        stripped = line.strip()
        if stripped.startswith("```"):
            in_code_block = not in_code_block
            index += 1
            continue
        if not stripped:
            pdf.ln(4)
            index += 1
            continue
        if stripped.startswith("#"):
            level = len(stripped) - len(stripped.lstrip("#"))
            title = stripped[level:].strip()
            size = {1: 18, 2: 15, 3: 13, 4: 12}.get(level, 11)
            pdf.set_font("Helvetica", style="B", size=size)
            pdf_multi_cell(pdf, 8, normalize_pdf_text(title))
            pdf.ln(1)
            pdf.set_font("Helvetica", size=11)
            index += 1
            continue
        if stripped.startswith("- "):
            pdf_multi_cell(pdf, 6, normalize_pdf_text(f"- {stripped[2:].strip()}"))
            index += 1
            continue
        if is_numbered_line(stripped):
            pdf_multi_cell(pdf, 6, normalize_pdf_text(stripped))
            index += 1
            continue
        if stripped.startswith("|"):
            pdf_multi_cell(pdf, 6, normalize_pdf_text(stripped.replace("|", "  ")))
            index += 1
            continue
        if in_code_block:
            pdf.set_font("Courier", size=9)
            pdf_multi_cell(pdf, 5, normalize_pdf_text(line))
            pdf.set_font("Helvetica", size=11)
            index += 1
            continue
        style = ""
        paragraph_lines = [stripped]
        if stripped.startswith("**") and stripped.endswith("**"):
            style = "B"
            index += 1
        else:
            index += 1
            while index < len(lines):
                candidate = lines[index].strip()
                if is_special_line(candidate):
                    break
                paragraph_lines.append(candidate)
                index += 1
        text = " ".join(paragraph_lines).strip("*")
        pdf.set_font("Helvetica", style=style, size=11)
        pdf_multi_cell(pdf, 6, normalize_pdf_text(text))
    pdf.output(str(output_path))


def parse_slides(markdown_path: Path) -> list[dict[str, list[str] | str]]:
    text = markdown_path.read_text(encoding="utf-8")
    chunks = [chunk.strip() for chunk in text.split("\n---\n") if chunk.strip()]
    slides: list[dict[str, list[str] | str]] = []
    for chunk in chunks:
        lines = [line.strip() for line in chunk.splitlines() if line.strip()]
        if not lines:
            continue
        title = "Slide"
        bullets: list[str] = []
        for line in lines:
            if line.startswith("%"):
                continue
            if line.startswith("#"):
                title = line.lstrip("#").strip()
            elif line.startswith("- "):
                bullets.append(line[2:].strip())
            else:
                bullets.append(line)
        slides.append({"title": title, "bullets": bullets})
    return slides


def render_pptx(markdown_path: Path, output_path: Path) -> None:
    presentation = Presentation()
    presentation.slide_width = PptInches(13.333)
    presentation.slide_height = PptInches(7.5)

    slides = parse_slides(markdown_path)
    for slide_data in slides:
        slide = presentation.slides.add_slide(presentation.slide_layouts[6])
        title_box = slide.shapes.add_textbox(PptInches(0.6), PptInches(0.4), PptInches(12.0), PptInches(0.8))
        title_frame = title_box.text_frame
        title_frame.clear()
        title_paragraph = title_frame.paragraphs[0]
        title_paragraph.text = str(slide_data["title"])
        title_paragraph.font.size = PptPt(28)
        title_paragraph.font.bold = True
        title_paragraph.alignment = PP_ALIGN.LEFT

        body_box = slide.shapes.add_textbox(PptInches(0.9), PptInches(1.6), PptInches(11.5), PptInches(5.3))
        body_frame = body_box.text_frame
        body_frame.word_wrap = True
        first = True
        for bullet in slide_data["bullets"]:
            if first:
                paragraph = body_frame.paragraphs[0]
                first = False
            else:
                paragraph = body_frame.add_paragraph()
            paragraph.text = str(bullet)
            paragraph.font.size = PptPt(20)
            paragraph.level = 0
            paragraph.alignment = PP_ALIGN.LEFT

    presentation.save(output_path)


def main() -> int:
    parser = argparse.ArgumentParser(description="Render report and presentation deliverables")
    parser.add_argument("--report-md", default="docs/Report.md")
    parser.add_argument("--presentation-md", default="docs/Presentation.md")
    parser.add_argument("--report-docx", default="Report.docx")
    parser.add_argument("--report-pdf", default="Report.pdf")
    parser.add_argument("--presentation-pptx", default="Presentation.pptx")
    args = parser.parse_args()

    render_docx(Path(args.report_md), Path(args.report_docx))
    render_pdf(Path(args.report_md), Path(args.report_pdf))
    render_pptx(Path(args.presentation_md), Path(args.presentation_pptx))
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
